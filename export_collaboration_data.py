import json
import os

from openpyxl import load_workbook


# ============================================================
# FILE PATHS
# ============================================================

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

EXCEL_FILE = os.path.join(
    BASE_DIR,
    "Gridlock_Utilities_Resource_and_Collaboration_Analysis.xlsx"
)

OUTPUT_DIR = os.path.join(
    BASE_DIR,
    "server",
    "data"
)

OUTPUT_FILE = os.path.join(
    OUTPUT_DIR,
    "collaboration-savings.json"
)


# ============================================================
# HELPERS
# ============================================================

def clean_number(value):
    """
    Convert spreadsheet numbers into JSON-safe Python numbers.
    """
    if value is None:
        return 0

    if isinstance(value, (int, float)):
        return value

    try:
        return float(value)
    except (TypeError, ValueError):
        return 0


def clean_text(value):
    if value is None:
        return ""

    return str(value).strip()


# ============================================================
# LOAD WORKBOOK
# ============================================================

print("Loading Gridlock collaboration workbook...")

workbook = load_workbook(
    EXCEL_FILE,
    data_only=True
)

sheet = workbook["Collaboration Savings"]


# ============================================================
# READ HEADERS
# ============================================================

headers = [
    cell.value
    for cell in sheet[1]
]

header_index = {
    header: index
    for index, header in enumerate(headers)
    if header is not None
}


def get_value(row, column_name):
    index = header_index[column_name]

    return row[index]


# ============================================================
# EXPORT OPPORTUNITIES
# ============================================================

opportunities = []


for excel_row in sheet.iter_rows(
    min_row=2,
    max_row=8,
    values_only=True
):

    overlap_id = clean_text(
        get_value(
            excel_row,
            "Overlap ID"
        )
    )

    if not overlap_id:
        continue


    opportunity = {

        # ----------------------------------------------------
        # IDENTIFICATION
        # ----------------------------------------------------

        "overlapId": overlap_id,

        "georgiaProjectId": clean_text(
            get_value(
                excel_row,
                "Utility A Project ID"
            )
        ),

        "georgiaProjectName": clean_text(
            get_value(
                excel_row,
                "Utility A Project"
            )
        ),

        "georgiaInServiceDate": clean_text(
            get_value(
                excel_row,
                "A In-Service Date"
            )
        ),

        "dominionProjectId": clean_text(
            get_value(
                excel_row,
                "Utility B Project ID"
            )
        ),

        "dominionProjectName": clean_text(
            get_value(
                excel_row,
                "Utility B Project"
            )
        ),

        "dominionInServiceDate": clean_text(
            get_value(
                excel_row,
                "B In-Service Date"
            )
        ),


        # ----------------------------------------------------
        # COORDINATION
        # ----------------------------------------------------

        "distanceMiles": clean_number(
            get_value(
                excel_row,
                "Distance (mi)"
            )
        ),

        "timeGapDays": clean_number(
            get_value(
                excel_row,
                "Time Gap (days)"
            )
        ),

        "opportunityType": clean_text(
            get_value(
                excel_row,
                "Opportunity Type"
            )
        ),


        # ----------------------------------------------------
        # BASE PROJECT COSTS
        # ----------------------------------------------------

        "projectCosts": {

            "georgiaProjectCost": clean_number(
                get_value(
                    excel_row,
                    "A Total Project Cost ($)"
                )
            ),

            "dominionProjectCost": clean_number(
                get_value(
                    excel_row,
                    "B Total Project Cost ($)"
                )
            ),

            "baselineSeparateCost": clean_number(
                get_value(
                    excel_row,
                    "Baseline Separate Cost ($)"
                )
            ),

        },


        # ----------------------------------------------------
        # EQUIPMENT COSTS
        # ----------------------------------------------------

        "equipmentCosts": {

            "georgia": clean_number(
                get_value(
                    excel_row,
                    "A Equipment Cost ($)"
                )
            ),

            "dominion": clean_number(
                get_value(
                    excel_row,
                    "B Equipment Cost ($)"
                )
            ),

        },


        # ----------------------------------------------------
        # LOGISTICS COSTS
        # ----------------------------------------------------

        "logisticsCosts": {

            "georgia": clean_number(
                get_value(
                    excel_row,
                    "A Logistics Cost ($)"
                )
            ),

            "dominion": clean_number(
                get_value(
                    excel_row,
                    "B Logistics Cost ($)"
                )
            ),

        },


        # ----------------------------------------------------
        # CONTRACT COSTS
        # ----------------------------------------------------

        "contractCosts": {

            "georgia": clean_number(
                get_value(
                    excel_row,
                    "A Contract Cost ($)"
                )
            ),

            "dominion": clean_number(
                get_value(
                    excel_row,
                    "B Contract Cost ($)"
                )
            ),

        },


        # ----------------------------------------------------
        # MATERIAL COSTS
        # ----------------------------------------------------

        "materialCosts": {

            "georgia": clean_number(
                get_value(
                    excel_row,
                    "A Material Cost ($)"
                )
            ),

            "dominion": clean_number(
                get_value(
                    excel_row,
                    "B Material Cost ($)"
                )
            ),

        },


        # ----------------------------------------------------
        # POTENTIALLY SHAREABLE EQUIPMENT
        # ----------------------------------------------------

        "sharedResources": {

            "bucketTrucks": clean_number(
                get_value(
                    excel_row,
                    "Shared Bucket Trucks"
                )
            ),

            "diggerDerricks": clean_number(
                get_value(
                    excel_row,
                    "Shared Digger Derricks"
                )
            ),

            "cranes": clean_number(
                get_value(
                    excel_row,
                    "Shared Cranes"
                )
            ),

            "excavators": clean_number(
                get_value(
                    excel_row,
                    "Shared Excavators"
                )
            ),

        },


        # ----------------------------------------------------
        # MODEL FACTORS
        # ----------------------------------------------------

        "modelFactors": {

            "distanceFactor": clean_number(
                get_value(
                    excel_row,
                    "Distance Factor"
                )
            ),

            "timeFactor": clean_number(
                get_value(
                    excel_row,
                    "Time Factor"
                )
            ),

        },


        # ----------------------------------------------------
        # SAVINGS BREAKDOWN
        # ----------------------------------------------------

        "savings": {

            "equipment": clean_number(
                get_value(
                    excel_row,
                    "Equipment Savings ($)"
                )
            ),

            "logistics": clean_number(
                get_value(
                    excel_row,
                    "Logistics Savings ($)"
                )
            ),

            "contracting": clean_number(
                get_value(
                    excel_row,
                    "Contract Savings ($)"
                )
            ),

            "bulkMaterials": clean_number(
                get_value(
                    excel_row,
                    "Bulk Material Savings ($)"
                )
            ),

            "totalPotentialSavings": clean_number(
                get_value(
                    excel_row,
                    "Estimated Potential Savings ($)"
                )
            ),

            "estimatedCollaborativeCost": clean_number(
                get_value(
                    excel_row,
                    "Estimated Collaborative Cost ($)"
                )
            ),

            "savingsPercent": (
                clean_number(
                    get_value(
                        excel_row,
                        "Savings %"
                    )
                )
                * 100
            ),

        },


        # ----------------------------------------------------
        # EXISTING MODEL RECOMMENDATION
        # ----------------------------------------------------

        "recommendation": clean_text(
            get_value(
                excel_row,
                "Recommendation"
            )
        ),

    }


    opportunities.append(
        opportunity
    )


# ============================================================
# FINAL JSON
# ============================================================

output = {

    "metadata": {

        "source":
            "Gridlock Utilities Resource and Collaboration Analysis",

        "modelType":
            "Synthetic prototype financial/resource model",

        "disclaimer":
            "Resource requirements and financial estimates are modeled prototype values and are not actual internal utility cost data.",

        "opportunityCount":
            len(opportunities),

    },

    "opportunities":
        opportunities,

}


# ============================================================
# WRITE FILE
# ============================================================

os.makedirs(
    OUTPUT_DIR,
    exist_ok=True
)


with open(
    OUTPUT_FILE,
    "w",
    encoding="utf-8"
) as file:

    json.dump(
        output,
        file,
        indent=2
    )


print()
print("========================================")
print("GRIDLOCK EXPORT COMPLETE")
print("========================================")
print()
print(
    f"Exported {len(opportunities)} opportunities."
)
print()
print(
    f"Output: {OUTPUT_FILE}"
)